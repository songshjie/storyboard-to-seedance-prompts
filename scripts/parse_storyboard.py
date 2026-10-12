#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""解析视频分镜表（xlsx / csv）为结构化 JSON，供生成 Seedance 提示词使用。

自动识别常见列：镜头号 / 时长 / 景别 / 运镜 / 画面内容 / 人物角色 / 道具 / 台词音效 / 导演备注；
按时长累加生成每镜起止时间码；如存在「人物明细」「道具明细」子表则一并抽取锚点清单。

用法:
  python parse_storyboard.py <input.xlsx|csv> [--output out.json] [--sheet 分镜脚本]

输出 JSON 结构:
  {
    "source": "...",
    "total_duration": 229.0,
    "characters": [{"name": "容辞", "desc": "..."}],
    "props": [{"name": "手机", "desc": "..."}],
    "shots": [
      {"no": 1, "start": 0.0, "end": 3.0, "duration": 3.0,
       "shot_size": "特写", "movement": "固定", "content": "...",
       "character": "容辞", "prop": "手机", "dialogue": "...", "note": "..."}
    ]
  }
"""

import argparse
import csv
import json
import os
import re
import sys

try:
    from openpyxl import load_workbook
    HAVE_OPENPYXL = True
except ImportError:  # pragma: no cover
    HAVE_OPENPYXL = False

COLUMN_ALIASES = {
    "shot_no": ["镜头号", "镜号", "镜头序号", "序号", "编号", "镜头", "shot", "shot no", "no"],
    "duration": ["时长", "时长(s)", "时长（秒）", "时长秒", "时间", "秒", "duration", "dur"],
    "shot_size": ["景别", "景别/构图", "shot size"],
    "movement": ["运镜", "运镜方式", "镜头运动", "运镜/镜头", "机位", "movement"],
    "content": ["画面内容", "画面描述", "画面", "内容", "描述", "镜头内容", "content"],
    "character": ["人物角色", "人物", "角色", "演员", "character"],
    "prop": ["道具", "道具明细", "prop"],
    "dialogue": ["台词/音效", "台词音效", "台词", "音效", "对白", "dialogue"],
    "note": ["导演备注", "备注", "说明", "note"],
}

CHARACTER_SHEET_ALIASES = ["人物明细", "人物", "角色明细", "角色", "characters"]
PROP_SHEET_ALIASES = ["道具明细", "道具", "props"]
STORYBOARD_SHEET_ALIASES = ["分镜脚本", "分镜", "镜头表", "脚本", "storyboard", "shots"]


def norm(s):
    if s is None:
        return ""
    return re.sub(r"\s+", "", str(s))


def to_float(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("秒", "").replace("s", "").replace("S", "")
    try:
        return float(s)
    except ValueError:
        m = re.search(r"\d+(\.\d+)?", s)
        return float(m.group(0)) if m else 0.0


def pick_storyboard_sheet(workbook, sheet_names, rows_by_sheet):
    for name in sheet_names:
        if norm(name) in [norm(a) for a in STORYBOARD_SHEET_ALIASES]:
            return name
    # 选行数最多且非空的表
    best, best_n = None, -1
    for name, rows in rows_by_sheet.items():
        n = sum(1 for r in rows if any(c is not None and str(c).strip() for c in r))
        if n > best_n:
            best, best_n = name, n
    return best


def detect_header_index(rows):
    best_idx, best_score = None, 0
    for i, row in enumerate(rows[:10]):
        cells = [norm(c) for c in row]
        score = 0
        for aliases in COLUMN_ALIASES.values():
            if any(c in [norm(a) for a in aliases] for c in cells):
                score += 1
        if score > best_score:
            best_idx, best_score = i, score
    return best_idx


def build_column_map(header_row):
    mapping = {}
    for field, aliases in COLUMN_ALIASES.items():
        for i, cell in enumerate(header_row):
            c = norm(cell).lower()
            if not c:
                continue
            if any(c == norm(a).lower() for a in aliases):
                mapping[field] = i
                break
    return mapping


def parse_storyboard_sheet(rows, header_idx, colmap):
    shots, cur = [], 0.0
    for row in rows[header_idx + 1:]:
        if not any(c is not None and str(c).strip() for c in row):
            continue
        def cell(field):
            i = colmap.get(field)
            if i is None or i >= len(row):
                return ""
            v = row[i]
            return "" if v is None else str(v).strip()

        duration = to_float(cell("duration"))
        if duration <= 0:
            duration = 0.0
        start, end = cur, cur + duration
        cur = end

        no = cell("shot_no")
        try:
            no = int(float(no))
        except (ValueError, TypeError):
            pass
        shots.append({
            "no": no,
            "start": round(start, 2),
            "end": round(end, 2),
            "duration": round(duration, 2),
            "shot_size": cell("shot_size"),
            "movement": cell("movement"),
            "content": cell("content"),
            "character": cell("character"),
            "prop": cell("prop"),
            "dialogue": cell("dialogue"),
            "note": cell("note"),
        })
    return shots, round(cur, 2)


def parse_people(rows):
    out = []
    header_idx = detect_header_index(rows)
    if header_idx is None:
        return out
    hdr = [norm(c) for c in rows[header_idx]]
    # 找“名字/人物”列与“描述/特征/形象”列
    name_col = desc_col = None
    for i, h in enumerate(hdr):
        hl = h.lower()
        if hl in ("名字", "姓名", "人物", "角色", "名称", "道具", "物品", "name", "character", "prop"):
            name_col = i
        elif any(k in hl for k in ("描述", "特征", "形象", "人设", "外形", "要点", "desc", "detail")):
            desc_col = i
    if name_col is None:
        return out
    for row in rows[header_idx + 1:]:
        if name_col >= len(row):
            continue
        v = row[name_col]
        if v is None or not str(v).strip():
            continue
        desc = ""
        if desc_col is not None and desc_col < len(row) and row[desc_col] is not None:
            desc = str(row[desc_col]).strip()
        out.append({"name": str(v).strip(), "desc": desc})
    return out


def load_xlsx(path, sheet_name=None):
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    rows_by_sheet = {}
    for name in sheet_names:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
        rows_by_sheet[name] = rows

    sb_name = sheet_name or pick_storyboard_sheet(wb, sheet_names, rows_by_sheet)
    sb_rows = rows_by_sheet.get(sb_name, [])
    header_idx = detect_header_index(sb_rows)
    if header_idx is None:
        raise ValueError(f"未在子表「{sb_name}」中找到表头")
    colmap = build_column_map(sb_rows[header_idx])
    shots, total = parse_storyboard_sheet(sb_rows, header_idx, colmap)

    characters, props = [], []
    for name in sheet_names:
        rows = rows_by_sheet[name]
        if norm(name) in [norm(a) for a in CHARACTER_SHEET_ALIASES]:
            characters = parse_people(rows)
        elif norm(name) in [norm(a) for a in PROP_SHEET_ALIASES]:
            props = parse_people(rows)
    return {
        "source": path,
        "sheet": sb_name,
        "columns_used": colmap,
        "total_duration": total,
        "characters": characters,
        "props": props,
        "shots": shots,
    }


def load_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.reader(f)]
    header_idx = detect_header_index(rows)
    if header_idx is None:
        raise ValueError("未找到表头行")
    colmap = build_column_map(rows[header_idx])
    shots, total = parse_storyboard_sheet(rows, header_idx, colmap)
    return {
        "source": path,
        "sheet": "csv",
        "columns_used": colmap,
        "total_duration": total,
        "characters": [],
        "props": [],
        "shots": shots,
    }


def main():
    ap = argparse.ArgumentParser(description="解析分镜表为结构化 JSON")
    ap.add_argument("input", help="分镜表文件路径（.xlsx / .csv）")
    ap.add_argument("--output", help="输出 JSON 路径（默认与输入同名 _parsed.json）")
    ap.add_argument("--sheet", help="指定分镜表所在子表名（默认自动探测）")
    args = ap.parse_args()

    if not os.path.exists(args.input):
        sys.exit(f"文件不存在: {args.input}")
    ext = os.path.splitext(args.input)[1].lower()
    if ext == ".xlsx":
        if not HAVE_OPENPYXL:
            sys.exit("缺少 openpyxl，请先 pip install openpyxl")
        data = load_xlsx(args.input, args.sheet)
    elif ext == ".csv":
        data = load_csv(args.input)
    else:
        sys.exit(f"不支持的文件类型: {ext}（仅支持 .xlsx / .csv）")

    out = args.output or os.path.splitext(args.input)[0] + "_parsed.json"
    with open(out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(json.dumps({
        "ok": True,
        "shots": len(data["shots"]),
        "total_duration": data["total_duration"],
        "sheet": data["sheet"],
        "columns_used": data["columns_used"],
        "characters": [c["name"] for c in data["characters"]],
        "props": [p["name"] for p in data["props"]],
        "output": out,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
